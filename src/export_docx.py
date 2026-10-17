"""
Generates academic Word document (.docx) from Phase 1 & 2 experimental results.
Strictly compliant with ISO/IEC 29500 (ECMA-376) OpenXML schema for 100% error-free loading in MS Word.
"""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


import lxml.etree


def set_cell_shading(cell, hex_color: str):
    """
    Applies background fill color to a table cell in strict compliance with ECMA-376 schema order.
    Removes any existing shd tag and places the new shd BEFORE vAlign, tcMar, etc.
    """
    tcPr = cell._tc.get_or_add_tcPr()

    # 1. Remove existing shd if already present
    for child in list(tcPr):
        if lxml.etree.QName(child).localname == "shd":
            tcPr.remove(child)

    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')

    # 2. In ECMA-376 tcPr, shd MUST come before noWrap, tcMar, textDirection, tcFitText, vAlign
    after_shd_tags = {
        "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign",
        "hideMark", "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"
    }

    inserted = False
    for child in tcPr:
        if lxml.etree.QName(child).localname in after_shd_tags:
            child.addprevious(shd)
            inserted = True
            break

    if not inserted:
        tcPr.append(shd)


def create_report_docx(output_path: Path):
    doc = docx.Document()

    # Standard Academic 1-inch margins (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document-wide Normal Style
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(6)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run(
        "BÁO CÁO NGHIÊN CỨU THỰC NGHIỆM VÀ KẾT QUẢ THỰC THI\n(PHASE 1 & PHASE 2)"
    )
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    # Subtitle / Metadata
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run(
        "Đề tài: Nghiên cứu và ứng dụng kỹ thuật Explainable AI (XAI) trong dự đoán nguy cơ mắc bệnh tim mạch\n"
        "và phân tích tác động của các chỉ số y sinh\n"
        "Học phần: Khóa luận Tốt nghiệp - Vai trò: Senior Machine Learning Engineer"
    )
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    def add_h1(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        return p

    def add_h2(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        return p

    def add_bullet(bold_prefix: str, text: str):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run_b = p.add_run(bold_prefix)
        run_b.font.bold = True
        run_b.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        p.add_run(text)
        return p

    def add_callout(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.right_indent = Inches(0.3)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        run.font.italic = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        return p

    def style_table(table, col_widths, align_cols):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        for row_idx, row in enumerate(table.rows):
            # Header repetition across pages
            if row_idx == 0:
                trPr = row._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for col_idx, cell in enumerate(row.cells):
                cell.width = col_widths[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.1

                align_type = align_cols.get(col_idx, WD_ALIGN_PARAGRAPH.LEFT)
                p.alignment = align_type

                run = p.runs[0] if p.runs else p.add_run()
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)

                if row_idx == 0:
                    set_cell_shading(cell, "1F4E79")
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif row_idx % 2 == 1:
                    set_cell_shading(cell, "F7F9FC")

    # ==========================================
    # PHẦN 1
    # ==========================================
    add_h1("PHẦN 1: QUY TRÌNH TIỀN XỬ LÝ & LÀM SẠCH DỮ LIỆU LÂM SÀNG (PHASE 1)")

    add_h2("1.1. Thống kê thực tế về quá trình sàng lọc và làm sạch dữ liệu")
    doc.add_paragraph(
        "Bộ dữ liệu nghiên cứu bệnh lý tim mạch (Cardiovascular Disease Dataset) chứa các bản ghi lâm sàng thu thập từ thực tế khám chữa bệnh. Dữ liệu y tế thô thường tồn tại nhiều sai số do lỗi nhập liệu của nhân viên y tế (human typographical errors), lỗi cảm biến đo huyết áp cơ học/điện tử, hoặc hiện tượng nghịch đảo huyết áp sinh lý. "
        "Dựa trên kết quả thực thi thực tế của Pipeline preprocess.py, số liệu thống kê thu được như sau:"
    )

    t1_data = [
        ["Chỉ số sàng lọc lâm sàng", "Số lượng bản ghi", "Tỷ lệ (%)", "Ghi chú kỹ thuật"],
        ["Tổng số bản ghi ban đầu", "70,000", "100.00%", "Toàn bộ dữ liệu thô đưa vào pipeline"],
        ["- Bản ghi sai lệch huyết áp tâm thu (ap_hi < 90 hoặc > 250)", "2,268", "3.24%", "Lỗi nhập liệu (-100, 16, 300, 14,000)"],
        ["- Bản ghi sai lệch huyết áp tâm trương (ap_lo < 60 hoặc > 150)", "2,500", "3.57%", "Lỗi nhập liệu (0, 20, 180, 1,100)"],
        ["- Bản ghi huyết áp nghịch đảo sinh lý (ap_hi <= ap_lo)", "1,795", "2.56%", "Áp lực tâm thu thấp hơn tâm trương"],
        ["- Bản ghi dị thường thể chất & BMI (chiều cao, cân nặng, BMI)", "436", "0.62%", "Chiều cao < 100cm, cân nặng < 30kg, BMI ngoài 10-80"],
        ["Tổng số bản ghi nhiễu bị loại bỏ", "4,802", "6.86%", "Hợp của các điều kiện lỗi vi phạm sinh lý"],
        ["Số bản ghi đạt chuẩn lâm sàng giữ lại", "65,198", "93.14%", "Tập dữ liệu sạch chất lượng cao"],
        ["- Tập huấn luyện (Train set - 80%)", "52,158", "74.51%", "Phân tầng Stratified (Tỷ lệ nhãn bệnh: 58.45%)"],
        ["- Tập kiểm định độc lập (Test set - 20%)", "13,040", "18.63%", "Phân tầng Stratified (Tỷ lệ nhãn bệnh: 58.45%)"],
    ]

    table1 = doc.add_table(rows=len(t1_data), cols=4)
    for r_idx, row in enumerate(table1.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t1_data[r_idx][c_idx]

    col_widths1 = [Inches(2.8), Inches(1.2), Inches(0.9), Inches(2.2)]
    align1 = {0: WD_ALIGN_PARAGRAPH.LEFT, 1: WD_ALIGN_PARAGRAPH.RIGHT, 2: WD_ALIGN_PARAGRAPH.RIGHT, 3: WD_ALIGN_PARAGRAPH.LEFT}
    style_table(table1, col_widths1, align1)

    # Highlight special summary rows
    for r_idx in [1, 6, 7]:
        for cell in table1.rows[r_idx].cells:
            set_cell_shading(cell, "EAF0F6")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    add_callout(
        "Ghi chú phân tích: Tổng số dòng bị loại bỏ (4,802 dòng) nhỏ hơn tổng đại số các lỗi đơn lẻ vì một bệnh nhân có thể đồng thời vi phạm nhiều tiêu chí (ví dụ: vừa có huyết áp tâm thu cao dị biệt vừa có huyết áp tâm trương sai lệch)."
    )

    add_h2("1.2. Cơ sở y khoa và nền tảng toán học của các bước biến đổi đặc trưng")
    add_bullet(
        "A. Quy đổi tuổi từ ngày sang năm (Age_years = Age_days / 365.25): ",
        "Trong y văn tim mạch quốc tế (AHA/ACC, ESC), các thang điểm đánh giá nguy cơ biến cố tim mạch 10 năm (như Framingham, SCORE2) đều sử dụng biến số tuổi theo năm. Chia cho 365.25 (tính chu kỳ năm nhuận) giúp đưa giá trị về dải tự nhiên [29.5, 64.9] tuổi, loại bỏ phương sai không cần thiết và giúp biểu đồ giải thích SHAP/LIME ở Phase 3 mang ý nghĩa trực quan với y bác sĩ."
    )
    add_bullet(
        "B. Kỹ thuật sinh đặc trưng Chỉ số khối cơ thể (BMI = Weight / Height^2): ",
        "Độc lập chiều cao hoặc cân nặng không phản ánh được mức độ béo phì. Một bệnh nhân nặng 80kg nhưng cao 1m85 có thể trạng bình thường, trong khi cùng cân nặng đó ở người cao 1m50 là béo phì độ II. BMI nén tương quan phi tuyến giữa chiều cao và cân nặng thành đại lượng chuẩn hóa, giúp thuật toán phân nhánh cây tìm được điểm cắt phân loại bệnh lý tương ứng với ngưỡng WHO (Bình thường < 25, Tiền béo phì 25-30, Béo phì >= 30)."
    )
    add_bullet(
        "C. Lọc bỏ các giá trị nghịch lý huyết áp (ap_hi <= ap_lo và ngoài [90-250]/[60-150]): ",
        "Về huyết động lực học, áp lực máu tâm thu luôn phải lớn hơn áp lực máu tâm trương để tạo chênh lệch áp suất tống máu (Pulse Pressure PP = ap_hi - ap_lo >= 20 mmHg). Bản ghi có ap_hi <= ap_lo là bất khả thi về mặt sinh học. Đồng thời, khoảng [90-250] mmHg và [60-150] mmHg đại diện cho dải sinh tồn ngoại trú."
    )
    add_bullet(
        "D. Chuẩn hóa dữ liệu bằng StandardScaler và Chống rò rỉ (Data Leakage): ",
        "StandardScaler đưa các biến liên tục về phân phối có trung bình 0 và độ lệch chuẩn 1. Phép tính kỳ vọng và độ lệch chuẩn được thực hiện nghiêm ngặt chỉ trên tập Train (52,158 dòng). Các đặc trưng nhị phân/thứ bậc (gender, smoke, alco, active, cholesterol, gluc) được giữ nguyên giá trị rời rạc để đảm bảo tính diễn giải y học cho XAI."
    )

    # ==========================================
    # PHẦN 2
    # ==========================================
    add_h1("PHẦN 2: CƠ CHẾ HOẠT ĐỘNG & HUẤN LUYỆN MÔ HÌNH (PHASE 2)")

    add_h2("2.1. Cơ chế toán học của 3 thuật toán phân loại nguy cơ tim mạch")
    add_bullet(
        "1. Extreme Gradient Boosting (XGBoost): ",
        "XGBoost xây dựng chuỗi các cây quyết định hồi quy tuần tự. Tại mỗi vòng lặp, mô hình tối ưu hóa xấp xỉ bậc 2 Taylor của hàm mất mát logloss với sự tham gia của đạo hàm bậc 1 (gradient g_i) và đạo hàm bậc 2 (hessian h_i). Khâu điều hòa phạt độ phức tạp cây (gamma*T + 0.5*lambda*w^2 + alpha*|w|) kiểm soát chống quá khớp, nắm bắt hoàn hảo các tương tác phi tuyến phức hợp giữa các chỉ số tim mạch."
    )
    add_bullet(
        "2. Random Forest: ",
        "Mô hình quần thể dựa trên nguyên lý Bagging (Bootstrap Aggregating) kết hợp ngẫu nhiên hóa không gian đặc trưng. Thuật toán lấy mẫu có hoàn lại từ tập Train và chỉ xét tập con căn bậc 2 số đặc trưng tại mỗi nút phân nhánh theo chỉ số Gini Impurity. Random Forest có khả năng giảm phương sai dự đoán vượt trội và chống chịu nhiễu tốt."
    )
    add_bullet(
        "3. Support Vector Machine (SVM): ",
        "Tìm kiếm siêu phẳng phân tách tối ưu với khoảng cách lề cực đại trong không gian đặc trưng nhiều chiều. Sử dụng hàm nhân RBF để giải quyết quan hệ phi tuyến. Tuy nhiên, SVM gặp nút thắt tính toán ma trận Gram và hiệu chuẩn xác suất Platt Scaling với độ phức tạp O(N^2) đến O(N^3), khiến thời gian huấn luyện và suy luận kéo dài đáng kể trên tập dữ liệu lớn."
    )

    add_h2("2.2. Quy trình tối ưu hóa siêu tham số bằng Optuna")
    doc.add_paragraph(
        "Nghiên cứu ứng dụng thuật toán Bayesian TPE (Tree-structured Parzen Estimator) của thư viện Optuna để mô hình hóa xác suất siêu tham số, tối đa hóa tỷ số cải thiện kỳ vọng (Expected Improvement). "
        "Quy trình được thực thi qua 15 trials với 5-Fold Stratified Cross-Validation trên tập Train (52,158 mẫu) nhằm tối đa hóa ROC-AUC."
    )

    t2_data = [
        ["Siêu tham số (Hyperparameter)", "Không gian tìm kiếm", "Loại", "Giá trị tối ưu tìm được", "Ý nghĩa kiểm soát mô hình"],
        ["n_estimators", "[100, 600], bước 50", "Integer", "400", "Số lượng cây quyết định đủ lớn để học đặc trưng"],
        ["max_depth", "[3, 9]", "Integer", "4", "Cây nông giúp tránh ghi nhớ mẫu nhiễu cá thể"],
        ["learning_rate", "[0.01, 0.2]", "Float (Log)", "0.0604", "Tốc độ học ổn định, hội tụ mượt mà"],
        ["subsample", "[0.6, 1.0], bước 0.1", "Float", "0.70", "Lấy mẫu 70% dữ liệu mỗi cây chống quá khớp"],
        ["colsample_bytree", "[0.6, 1.0], bước 0.1", "Float", "0.60", "Lấy mẫu 60% đặc trưng mỗi cây chống phụ thuộc cục bộ"],
        ["min_child_weight", "[1, 10]", "Integer", "10", "Ràng buộc trọng số lá tối thiểu"],
        ["gamma", "[0.0, 5.0], bước 0.5", "Float", "4.0", "Ngưỡng phạt tạo nhánh mới nếu không giảm loss đáng kể"],
        ["reg_alpha (L1)", "[0.001, 10.0]", "Float (Log)", "8.2483", "Triệt tiêu trọng số yếu, tạo tính thưa (Sparsity)"],
        ["reg_lambda (L2)", "[0.001, 10.0]", "Float (Log)", "0.0018", "Điều hòa mượt trọng số lá cây"],
    ]

    table2 = doc.add_table(rows=len(t2_data), cols=5)
    for r_idx, row in enumerate(table2.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t2_data[r_idx][c_idx]

    col_widths2 = [Inches(1.8), Inches(1.3), Inches(0.8), Inches(1.1), Inches(2.1)]
    align2 = {0: WD_ALIGN_PARAGRAPH.LEFT, 1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.CENTER, 3: WD_ALIGN_PARAGRAPH.CENTER, 4: WD_ALIGN_PARAGRAPH.LEFT}
    style_table(table2, col_widths2, align2)

    add_callout(
        "Kết quả đánh giá chéo: Điểm ROC-AUC trung bình qua 5-Fold Stratified Cross-Validation trên tập Train đạt 0.8020, chứng minh cấu hình siêu tham số có độ ổn định và khái quát hóa cao."
    )

    # ==========================================
    # PHẦN 3
    # ==========================================
    add_h1("PHẦN 3: BẢNG ĐÁNH GIÁ & SO SÁNH THỰC NGHIỆM TRÊN TẬP TEST")

    add_h2("3.1. Bảng tổng hợp các chỉ số thực tế trên tập kiểm định độc lập (N = 13,040)")
    doc.add_paragraph(
        "Dưới đây là bảng kết quả đo lường khách quan trên tập Test (13,040 bản ghi chưa từng xuất hiện trong quá trình huấn luyện hay tinh chỉnh siêu tham số):"
    )

    t3_data = [
        ["Hạng", "Mô hình (Model)", "ROC-AUC", "Recall (Độ nhạy)", "Precision (Độ CX)", "F1-Score", "Accuracy", "Thời gian suy luận"],
        ["1", "XGBoost (Optuna Tuned)", "0.8009", "0.8099", "0.7517", "0.7797", "0.7325", "0.0150 s"],
        ["2", "XGBoost (Baseline)", "0.7996", "0.8052", "0.7505", "0.7769", "0.7297", "0.0150 s"],
        ["3", "Random Forest", "0.7979", "0.8129", "0.7470", "0.7785", "0.7297", "0.0869 s"],
        ["4", "Support Vector Machine (SVM)", "0.7856", "0.8208", "0.7444", "0.7807", "0.7305", "4.4599 s"],
    ]

    table3 = doc.add_table(rows=len(t3_data), cols=8)
    for r_idx, row in enumerate(table3.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = t3_data[r_idx][c_idx]

    col_widths3 = [Inches(0.6), Inches(2.0), Inches(0.9), Inches(1.0), Inches(1.0), Inches(0.8), Inches(0.8), Inches(1.0)]
    align3 = {0: WD_ALIGN_PARAGRAPH.CENTER, 1: WD_ALIGN_PARAGRAPH.LEFT, 2: WD_ALIGN_PARAGRAPH.RIGHT, 3: WD_ALIGN_PARAGRAPH.RIGHT, 4: WD_ALIGN_PARAGRAPH.RIGHT, 5: WD_ALIGN_PARAGRAPH.RIGHT, 6: WD_ALIGN_PARAGRAPH.RIGHT, 7: WD_ALIGN_PARAGRAPH.RIGHT}
    style_table(table3, col_widths3, align3)

    # Highlight winner row
    for cell in table3.rows[1].cells:
        set_cell_shading(cell, "EAF0F6")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    add_h2("3.2. Phân tích lý do vượt trội của mô hình chiến thắng (XGBoost Tuned)")
    add_bullet(
        "1. Năng lực phân loại tổng quát dẫn đầu (ROC-AUC = 0.8009): ",
        "ROC-AUC đo lường khả năng xếp hạng rủi ro trên mọi ngưỡng phân loại xác suất. XGBoost đạt 0.8009, vượt trội hơn Random Forest (0.7979) và SVM (0.7856), thể hiện tính nhất quán và độ tin cậy phân tầng bệnh nhân cao nhất."
    )
    add_bullet(
        "2. Độ nhạy lâm sàng cao (Recall = 80.99%) kiểm soát tốt Precision (75.17%): ",
        "Trong y học dự phòng tim mạch, chỉ số Recall (Độ nhạy - Sensitivity) là ưu tiên số một nhằm hạn chế tối đa tình trạng âm tính giả (bỏ sót bệnh nhân nguy hiểm). XGBoost phát hiện chính xác gần 81/100 ca có nguy cơ thực tế, đồng thời kiểm soát Precision ở mức 75.17%, tránh lãng phí nguồn lực xét nghiệm chuyên sâu."
    )
    add_bullet(
        "3. Tối ưu hóa điều hòa và loại bỏ quá khớp (Overfitting): ",
        "Nhờ siêu tham số reg_alpha = 8.2483 và max_depth = 4, mô hình triệt tiêu các đặc trưng gây nhiễu, thu hẹp khoảng cách giữa ROC-AUC tập Train CV (0.8020) và tập Test (0.8009) xuống chỉ 0.0011."
    )
    add_bullet(
        "4. Tốc độ suy luận thời gian thực (0.015 giây): ",
        "XGBoost xử lý toàn bộ 13,040 bệnh nhân chỉ trong 0.015s (nhanh hơn gần 300 lần so với SVM - 4.46s). Đây là cơ sở kỹ thuật bắt buộc để tích hợp vào Dashboard lâm sàng ở Phase 4 phục vụ bác sĩ trực tiếp mà không gây trễ hệ thống."
    )

    add_h2("3.3. Cơ sở khoa học để lựa chọn mô hình cho Phase 3 (XAI Integration)")
    doc.add_paragraph(
        "Mô hình XGBoost (Optuna Tuned) chính thức được lựa chọn làm Champion Model nhờ khả năng tương thích bản địa tuyệt vời với thuật toán TreeSHAP (Lundberg et al., 2020). "
        "TreeSHAP có thể tính toán chính xác giá trị đóng góp biên Shapley trong thời gian đa thức O(TLD^2), vượt trội hoàn toàn so với KernelSHAP dạng black-box của SVM, tạo tiền đề vững chắc để xuất các biểu đồ SHAP Beeswarm, Dependence, Waterfall và LIME ở Phase 3."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Academic report saved successfully to: {output_path}")


if __name__ == "__main__":
    out_file = Path(__file__).resolve().parent.parent / "docs" / "Bao_Cao_Phase_1_2.docx"
    create_report_docx(out_file)
